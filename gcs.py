# ============================================================
# SECTION 0 — IMPORTS
# ============================================================
import os
import re
import uuid
import logging
import time
import asyncio
import math
import secrets
import hmac
import hashlib
import base64
from dataclasses import dataclass
from typing import Literal, Optional
from pathlib import Path
from threading import Lock
from datetime import datetime, timezone, timedelta
from zoneinfo import ZoneInfo
from contextlib import asynccontextmanager
from contextvars import ContextVar
from io import BytesIO
import html
import json
from urllib.parse import quote

import pandas as pd
from pydantic import BaseModel, Field, ConfigDict
from openpyxl import Workbook
from fastapi import FastAPI, Request, Header, HTTPException, Cookie
from fastapi.responses import FileResponse, HTMLResponse, Response, StreamingResponse
from starlette.responses import JSONResponse
from starlette.middleware.base import BaseHTTPMiddleware
from google.cloud import bigquery, firestore, storage
from mcp.server.fastmcp import FastMCP
from fastapi.middleware.cors import CORSMiddleware

from graph_mcp import run_visualization_code
from mapping import (
    TABLE_TO_REPORT_NAMES,
    TABLE_TO_BQ_TABLE_NAME,
    ALLOWED_OPERATORS,
    USER_REGEX,
    BQ_TABLE_REGEX,
    SENSITIVE_COLUMNS,
)

# ============================================================
# SECTION 1 — COST CALCULATOR
# ============================================================
BQ_PRICE_PER_TB_USD: float = 6.25
CLOUD_RUN_REQUEST_PRICE_PER_M: float = 0.40
CLOUD_RUN_CPU_PRICE_PER_VCPU_SEC: float = 0.00002400
CLOUD_RUN_MEM_PRICE_PER_GIB_SEC: float = 0.00000250

EMAIL_REGEX = re.compile(r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$')
SAFE_COLUMN_REGEX = re.compile(r"^[A-Za-z0-9_]{1,128}$")
DOWNLOAD_FILENAME_REGEX = re.compile(
    r"^(?:chart_)?[A-Za-z0-9_]{1,64}_[a-f0-9]{8}_\d{8}_\d{6}\.(xlsx|png)$"
)
UNRESOLVED_TEMPLATE_REGEX = re.compile(
    r'(\$\{[^}]+\}|\{\{[^}]+\}\}|\{[^}]+\}|\[.*?MASKED.*?\])',
    re.IGNORECASE
)


@dataclass
class BigQueryCost:
    bytes_billed: int
    tb_billed: float
    cost_usd: float
    cost_thb: float


@dataclass
class CloudRunCost:
    duration_seconds: float
    vcpu: float
    memory_gib: float
    request_cost_usd: float
    cpu_cost_usd: float
    memory_cost_usd: float
    total_cost_usd: float
    total_cost_thb: float


@dataclass
class TotalCost:
    bigquery: BigQueryCost
    cloud_run: CloudRunCost
    grand_total_usd: float
    grand_total_thb: float
    usd_to_thb_rate: float
    summary: str


def calculate_bigquery_cost(bytes_billed, price_per_tb_usd=BQ_PRICE_PER_TB_USD, usd_to_thb=32.57):
    safe_bytes = max(int(bytes_billed), 0)
    tb_billed = safe_bytes / (1024 ** 4)
    cost_usd = tb_billed * price_per_tb_usd
    cost_thb = cost_usd * usd_to_thb
    return BigQueryCost(safe_bytes, round(tb_billed, 10), round(cost_usd, 8), round(cost_thb, 6))


def calculate_cloud_run_cost(duration_seconds, vcpu=1.0, memory_gib=0.5,
                              request_price_per_million=CLOUD_RUN_REQUEST_PRICE_PER_M,
                              cpu_price_per_vcpu_sec=CLOUD_RUN_CPU_PRICE_PER_VCPU_SEC,
                              mem_price_per_gib_sec=CLOUD_RUN_MEM_PRICE_PER_GIB_SEC,
                              usd_to_thb=32.57):
    billable_seconds = max(math.ceil(duration_seconds * 10) / 10, 0.1)
    request_cost_usd = request_price_per_million / 1_000_000
    cpu_cost_usd = vcpu * billable_seconds * cpu_price_per_vcpu_sec
    memory_cost_usd = memory_gib * billable_seconds * mem_price_per_gib_sec
    total_usd = request_cost_usd + cpu_cost_usd + memory_cost_usd
    total_thb = total_usd * usd_to_thb
    return CloudRunCost(
        round(billable_seconds, 2), vcpu, memory_gib,
        round(request_cost_usd, 10), round(cpu_cost_usd, 8), round(memory_cost_usd, 8),
        round(total_usd, 8), round(total_thb, 6),
    )


def calculate_total_cost(bytes_billed, duration_seconds, vcpu=1.0, memory_gib=0.5,
                          usd_to_thb=32.57, bq_price_per_tb_usd=BQ_PRICE_PER_TB_USD):
    bq = calculate_bigquery_cost(bytes_billed, bq_price_per_tb_usd, usd_to_thb)
    cr = calculate_cloud_run_cost(duration_seconds, vcpu, memory_gib, usd_to_thb=usd_to_thb)
    grand_usd = bq.cost_usd + cr.total_cost_usd
    grand_thb = grand_usd * usd_to_thb
    summary = (
        f"BQ: ${bq.cost_usd:.6f} ({bq.tb_billed:.6f} TB) | "
        f"Cloud Run: ${cr.total_cost_usd:.6f} ({cr.duration_seconds}s) | "
        f"Total: ${grand_usd:.6f} / ฿{grand_thb:.4f}"
    )
    return TotalCost(bq, cr, round(grand_usd, 8), round(grand_thb, 6), usd_to_thb, summary)


def format_cost_breakdown(cost):
    return {
        "bigquery": {
            "bytes_billed": cost.bigquery.bytes_billed,
            "tb_billed": cost.bigquery.tb_billed,
            "cost_usd": cost.bigquery.cost_usd,
            "cost_thb": cost.bigquery.cost_thb,
        },
        "cloud_run": {
            "duration_seconds": cost.cloud_run.duration_seconds,
            "vcpu": cost.cloud_run.vcpu,
            "memory_gib": cost.cloud_run.memory_gib,
            "request_cost_usd": cost.cloud_run.request_cost_usd,
            "cpu_cost_usd": cost.cloud_run.cpu_cost_usd,
            "memory_cost_usd": cost.cloud_run.memory_cost_usd,
            "total_cost_usd": cost.cloud_run.total_cost_usd,
            "total_cost_thb": cost.cloud_run.total_cost_thb,
        },
        "grand_total_usd": cost.grand_total_usd,
        "grand_total_thb": cost.grand_total_thb,
        "usd_to_thb_rate": cost.usd_to_thb_rate,
        "summary": cost.summary,
    }

def _normalize_gcs_prefix(prefix: str) -> str:
    cleaned = (prefix or "reports/").strip().strip("/")
    return f"{cleaned}/" if cleaned else "reports/"


def _safe_report_filename(file_name: str) -> str:
    if not DOWNLOAD_FILENAME_REGEX.fullmatch(file_name or ""):
        raise HTTPException(status_code=400, detail="รูปแบบชื่อไฟล์ไม่ถูกต้อง")
    return file_name


def _gcs_blob_name(file_name: str) -> str:
    safe_name = _safe_report_filename(file_name)
    return f"{_normalize_gcs_prefix(GCS_REPORT_PREFIX)}{safe_name}"


# ============================================================
# SECTION 2 — CONFIGURATION
# ============================================================
BANGKOK_TZ = ZoneInfo("Asia/Bangkok")
PROJECT_ID = os.environ.get("PROJECT_ID", "sc-ai-uat")
DATASET_ID = os.environ.get("DATASET_ID", "SCReport")
AUTH_TABLE = os.environ.get("AUTH_TABLE", "AuthenByMenu")
LOG_TABLE = os.environ.get("LOG_TABLE", "AiLog")
DEFAULT_USERNAME = ""
PUBLIC_BASE_URL = os.environ.get(
    "PUBLIC_BASE_URL",
    "https://sc-report-866803019306.asia-southeast3.run.app",
)

COST_USD_TO_THB = float(os.environ.get("COST_USD_TO_THB", "32.57"))
COST_CLOUD_RUN_VCPU = float(os.environ.get("COST_CLOUD_RUN_VCPU", "1.0"))
COST_CLOUD_RUN_MEM_GIB = float(os.environ.get("COST_CLOUD_RUN_MEM_GIB", "0.5"))
COST_BQ_PRICE_PER_TB = float(os.environ.get("COST_BQ_PRICE_PER_TB", "6.25"))

CLEANUP_SECRET = os.environ.get("CLEANUP_SECRET")
LIBRECHAT_JWT_SECRET = os.environ.get("LIBRECHAT_JWT_SECRET", "")

# ── GCS ──────────────────────────────────────────────────────
REPORT_BUCKET = os.environ.get("REPORT_BUCKET", "sc-report-files")
GCS_REPORT_PREFIX = os.environ.get("GCS_REPORT_PREFIX", "reports/")
# signed URL TTL (วินาที) — default 24 ชม.
GCS_SIGNED_URL_TTL = int(os.environ.get("GCS_SIGNED_URL_TTL", str(24 * 3600)))
# ─────────────────────────────────────────────────────────────

SESSION_COOKIE_NAME = "sc_report_session"
SESSION_MAX_AGE = 3600

# ============================================================
# SECTION 3 — CLIENTS & CACHES
# ============================================================
bq_client = bigquery.Client(project=PROJECT_ID)
firestore_client = firestore.Client(project=PROJECT_ID)
gcs_client = storage.Client(project=PROJECT_ID)          # ← NEW
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("SCReportSecurity")

# /tmp ยังคงใช้เป็น staging area ชั่วคราวก่อนอัปโหลด GCS
REPORT_DIR = Path("/tmp/reports").resolve()
REPORT_DIR.mkdir(parents=True, exist_ok=True)

_MAPPING_CACHE: dict[str, tuple[dict[str, str], datetime]] = {}
_MAPPING_CACHE_LOCK = Lock()
CACHE_TTL = timedelta(minutes=5)

_SCHEMA_CACHE: dict[str, tuple[dict[str, str], datetime]] = {}
_SCHEMA_CACHE_LOCK = Lock()
SCHEMA_CACHE_TTL = timedelta(hours=1)

_USERNAME_CACHE: dict[str, tuple[Optional[str], datetime]] = {}
_USERNAME_CACHE_LOCK = Lock()
USERNAME_CACHE_TTL = timedelta(minutes=10)

CURRENT_REQUEST_USERNAME: ContextVar[Optional[str]] = ContextVar(
    "current_request_username", default=None
)
CURRENT_REQUEST_BQ_USERNAME: ContextVar[Optional[str]] = ContextVar(
    "current_request_bq_username", default=None
)
CURRENT_REQUEST_BASE_URL: ContextVar[Optional[str]] = ContextVar(
    "current_request_base_url", default=None
)


# ============================================================
# SECTION 4 — GCS HELPERS
# ============================================================
def _normalize_gcs_prefix(prefix: str) -> str:
    cleaned = (prefix or "reports/").strip().strip("/")
    return f"{cleaned}/" if cleaned else "reports/"


def _safe_report_filename(file_name: str) -> str:
    if not DOWNLOAD_FILENAME_REGEX.fullmatch(file_name or ""):
        raise HTTPException(status_code=400, detail="รูปแบบชื่อไฟล์ไม่ถูกต้อง")
    return file_name


def _gcs_blob_name(file_name: str) -> str:
    safe_name = _safe_report_filename(file_name)
    return f"{_normalize_gcs_prefix(GCS_REPORT_PREFIX)}{safe_name}"


def get_gcs_content_type(file_name: str) -> str:
    safe_name = _safe_report_filename(file_name)
    if safe_name.endswith(".png"):
        return "image/png"
    return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def get_gcs_blob(file_name: str):
    safe_name = _safe_report_filename(file_name)
    bucket = gcs_client.bucket(REPORT_BUCKET)
    blob = bucket.blob(_gcs_blob_name(safe_name))

    if not blob.exists():
        raise HTTPException(
            status_code=404,
            detail="ไฟล์หมดอายุหรือไม่พบในระบบ",
        )

    return blob


def upload_to_gcs(local_path: Path, file_name: str) -> None:
    safe_name = _safe_report_filename(file_name)

    bucket = gcs_client.bucket(REPORT_BUCKET)
    blob = bucket.blob(_gcs_blob_name(safe_name))

    blob.upload_from_filename(
        str(local_path),
        content_type=get_gcs_content_type(safe_name),
    )

    logger.info(
        "[GCS] Uploaded '%s' -> gs://%s/%s",
        safe_name,
        REPORT_BUCKET,
        blob.name,
    )

    try:
        local_path.unlink(missing_ok=True)
    except Exception:
        logger.warning(
            "[GCS] Failed to remove local staging file: %s",
            local_path,
            exc_info=True,
        )


def generate_gcs_signed_url(file_name: str) -> str:
    safe_name = _safe_report_filename(file_name)
    blob = get_gcs_blob(safe_name)

    return blob.generate_signed_url(
        version="v4",
        expiration=timedelta(seconds=GCS_SIGNED_URL_TTL),
        method="GET",
        response_disposition=f'attachment; filename="{safe_name}"',
        response_type=get_gcs_content_type(safe_name),
    )


def download_from_gcs(file_name: str) -> tuple[bytes, str]:
    """
    Backward-compatible helper.
    Prefer StreamingResponse in /download for large files.
    """
    safe_name = _safe_report_filename(file_name)
    blob = get_gcs_blob(safe_name)
    return blob.download_as_bytes(), get_gcs_content_type(safe_name)


def list_gcs_files() -> list[dict]:
    bucket = gcs_client.bucket(REPORT_BUCKET)
    gcs_prefix = _normalize_gcs_prefix(GCS_REPORT_PREFIX)

    files = []
    for blob in bucket.list_blobs(prefix=gcs_prefix):
        name = blob.name.removeprefix(gcs_prefix)

        if not name:
            continue

        if not DOWNLOAD_FILENAME_REGEX.fullmatch(name):
            logger.warning(
                "[GCS] Ignoring unexpected object under report prefix: %s",
                blob.name,
            )
            continue

        files.append({
            "name": name,
            "size_bytes": blob.size or 0,
            "size_mb": round((blob.size or 0) / (1024 * 1024), 4),
            "created_at": blob.time_created.isoformat() if blob.time_created else "",
            "modified_at": blob.updated.isoformat() if blob.updated else "",
        })

    return sorted(files, key=lambda x: x["modified_at"], reverse=True)


def delete_gcs_file(file_name: str) -> None:
    safe_name = _safe_report_filename(file_name)
    blob = get_gcs_blob(safe_name)
    blob.delete()

    logger.info("[GCS] Deleted '%s'", safe_name)


def cleanup_gcs_files() -> int:
    cleanup_strategy = os.environ.get("CLEANUP_STRATEGY", "age").lower()
    cleanup_age_seconds = int(os.environ.get("CLEANUP_AGE_SECONDS", "86400"))
    cleanup_max_files = int(os.environ.get("CLEANUP_MAX_FILES", "20"))

    files = list_gcs_files()
    deleted = 0
    now = datetime.now(timezone.utc)

    def delete_file(name: str) -> None:
        nonlocal deleted
        get_gcs_blob(name).delete()
        deleted += 1

    if cleanup_strategy == "age":
        cutoff = now - timedelta(seconds=cleanup_age_seconds)

        for f in files:
            raw_modified_at = f.get("modified_at") or ""

            try:
                modified_at = datetime.fromisoformat(raw_modified_at)
            except ValueError:
                logger.warning(
                    "[GCS] Skipping cleanup for invalid modified_at: %s",
                    f,
                )
                continue

            if modified_at.tzinfo is None:
                modified_at = modified_at.replace(tzinfo=timezone.utc)

            if modified_at.astimezone(timezone.utc) < cutoff:
                delete_file(f["name"])

    elif cleanup_strategy == "keep_latest":
        keep_count = max(cleanup_max_files, 0)

        for f in files[keep_count:]:
            delete_file(f["name"])

    elif cleanup_strategy == "aggressive":
        for f in files:
            delete_file(f["name"])

    else:
        logger.warning(
            "[GCS] Unknown CLEANUP_STRATEGY='%s'; no files deleted.",
            cleanup_strategy,
        )

    logger.info("[GCS] Cleanup completed - %s file(s) removed.", deleted)
    return deleted


# ============================================================
# SECTION 5 — VALIDATORS & HELPERS
# ============================================================
def is_valid_email(value: object) -> bool:
    if not value or not isinstance(value, str):
        return False
    val_strip = value.strip()
    if len(val_strip) > 254:
        return False
    if UNRESOLVED_TEMPLATE_REGEX.search(val_strip):
        return False
    return bool(EMAIL_REGEX.match(val_strip))


def sanitize_column_name(column: str) -> str:
    if not column or not SAFE_COLUMN_REGEX.match(column):
        raise PermissionError(f"คอลัมน์ '{column}' มีรูปแบบไม่ถูกต้องหรือไม่ปลอดภัย")
    return column


def mask_username(username):
    if not username:
        return "Anonymous"
    if "@" in username:
        name, domain = username.split("@", 1)
        if len(name) == 0:
            masked_name = "***"
        elif len(name) <= 2:
            masked_name = name[0] + "***"
        else:
            masked_name = name[0] + "***" + name[-1]
        return f"{masked_name}@{domain}"
    if len(username) == 0:
        return "***"
    if len(username) <= 2:
        return username[0] + "***"
    return username[0] + "***" + username[-1]


def clean_table_id(raw):
    if not raw:
        return ""
    return raw.strip().strip("`").strip()


def validate_table_id(table_id):
    if not table_id:
        return False
    full_bq_regex = re.compile(r"^[A-Za-z0-9_-]+\.[A-Za-z0-9_]+\.[A-Za-z0-9_]+$")
    return bool(BQ_TABLE_REGEX.match(table_id)) or bool(full_bq_regex.match(table_id))


def extract_short_table_name(table_id):
    clean = table_id.strip().strip("`").strip()
    return clean.split(".")[-1].lower()


def is_valid_user(user):
    if not user:
        return False
    return bool(USER_REGEX.match(user.strip()))


def is_sensitive_column(column):
    return bool(column and column.lower() in SENSITIVE_COLUMNS)


# ============================================================
# SECTION 6 — CACHES
# ============================================================
def get_table_column_mapping(table_id):
    cleaned_table = clean_table_id(table_id)
    short_name = extract_short_table_name(cleaned_table)
    bq_table_name = TABLE_TO_BQ_TABLE_NAME.get(short_name, cleaned_table)

    now = datetime.now(BANGKOK_TZ)
    with _MAPPING_CACHE_LOCK:
        cached = _MAPPING_CACHE.get(bq_table_name)
        if cached:
            mapping, expiry = cached
            if now < expiry:
                return mapping

    try:
        doc_ref = (
            firestore_client.collection("data_dictionary")
            .document(DATASET_ID)
            .collection("tables")
            .document(bq_table_name)
        )
        doc = doc_ref.get()
        mapping = {}
        if doc.exists:
            data = doc.to_dict() or {}
            columns_list = data.get("columns", [])
            for col in columns_list:
                name = col.get("name")
                desc = col.get("description")
                if name:
                    mapping[name] = desc if desc else name
            logger.info(f"[Firestore] Loaded mapping for '{bq_table_name}' with {len(mapping)} columns.")
        else:
            logger.warning(f"[Firestore] Document data_dictionary/{DATASET_ID}/tables/{bq_table_name} does not exist.")

        with _MAPPING_CACHE_LOCK:
            _MAPPING_CACHE[bq_table_name] = (mapping, now + CACHE_TTL)
        return mapping
    except Exception as e:
        logger.error(f"Error fetching column mapping from Firestore: {e}", exc_info=True)
        return {}


def get_table_schema(table_ref):
    now = datetime.now(BANGKOK_TZ)
    with _SCHEMA_CACHE_LOCK:
        cached = _SCHEMA_CACHE.get(table_ref)
        if cached:
            schema, expiry = cached
            if now < expiry:
                return schema
    try:
        table = bq_client.get_table(table_ref)
        schema = {field.name.lower(): field.field_type for field in table.schema}
        with _SCHEMA_CACHE_LOCK:
            _SCHEMA_CACHE[table_ref] = (schema, now + SCHEMA_CACHE_TTL)
        return schema
    except Exception as e:
        logger.error(f"[Schema] Failed to fetch schema for {table_ref}: {e}")
        return {}


# ============================================================
# SECTION 7 — LOG ID GENERATOR
# ============================================================
def generate_ai_log_id():
    return uuid.uuid4().int & 0x7FFFFFFFFFFFFFFF


# ============================================================
# SECTION 8 — JWT & USERNAME LOOKUP
# ============================================================
def extract_email_from_jwt(token):
    if not token:
        return None
    try:
        import jwt as pyjwt
        if LIBRECHAT_JWT_SECRET:
            payload = pyjwt.decode(token, LIBRECHAT_JWT_SECRET, algorithms=["HS256", "HS384", "HS512"])
        else:
            payload = pyjwt.decode(token, options={"verify_signature": False},
                                   algorithms=["HS256", "HS384", "HS512", "RS256"])

        email = payload.get("email") or payload.get("user_email") or payload.get("preferred_username")
        if not email:
            sub = payload.get("sub", "")
            if sub and "@" in sub:
                email = sub

        if email and is_valid_email(email):
            logger.info(f"[JWT] Extracted email from token: {mask_username(email)}")
            return email.strip().lower()
        return None
    except Exception as e:
        logger.warning(f"[JWT] Failed to decode token: {type(e).__name__}: {e}")
        return None


def _lookup_username_sync(email):
    if not is_valid_email(email):
        logger.warning(f"[lookup_username] Invalid email skipped: '{email}'")
        return None

    email_lower = email.strip().lower()
    now = datetime.now(BANGKOK_TZ)

    with _USERNAME_CACHE_LOCK:
        cached = _USERNAME_CACHE.get(email_lower)
        if cached:
            cached_username, expiry = cached
            if now < expiry:
                logger.info(f"[UsernameCache] HIT for {mask_username(email)}: {cached_username}")
                return cached_username

    try:
        query = f"""
            SELECT UserName
            FROM `{PROJECT_ID}.{DATASET_ID}.{AUTH_TABLE}`
            WHERE LOWER(Email) = @email
            LIMIT 1
        """
        job_config = bigquery.QueryJobConfig(
            query_parameters=[bigquery.ScalarQueryParameter("email", "STRING", email_lower)]
        )
        results = list(bq_client.query(query, job_config=job_config).result())
        username = results[0]["UserName"] if results else None

        with _USERNAME_CACHE_LOCK:
            _USERNAME_CACHE[email_lower] = (username, now + USERNAME_CACHE_TTL)

        if username:
            logger.info(f"[UsernameCache] MISS → resolved '{mask_username(email)}' to '{username}'")
        else:
            logger.warning(f"[UsernameCache] MISS → no UserName for '{mask_username(email)}'")
        return username
    except Exception as e:
        logger.error(f"[lookup_username] BQ error: {e}", exc_info=True)
        return None


async def lookup_username_async(email):
    return await asyncio.to_thread(_lookup_username_sync, email)


def lookup_username_from_email(email):
    return _lookup_username_sync(email)


# ============================================================
# SECTION 9 — SESSION TOKEN
# ============================================================
def _make_session_token():
    if not CLEANUP_SECRET:
        raise RuntimeError("CLEANUP_SECRET not configured")
    nonce = secrets.token_urlsafe(16)
    timestamp = str(int(time.time()))
    payload = f"{nonce}:{timestamp}"
    signature = hmac.new(CLEANUP_SECRET.encode(), payload.encode(), hashlib.sha256).digest()
    sig_b64 = base64.urlsafe_b64encode(signature).decode().rstrip("=")
    return f"{payload}:{sig_b64}"


def _verify_session_token(token):
    if not CLEANUP_SECRET or not token:
        return False
    try:
        parts = token.split(":")
        if len(parts) != 3:
            return False
        nonce, timestamp, sig_b64 = parts
        if int(time.time()) - int(timestamp) > SESSION_MAX_AGE:
            return False
        payload = f"{nonce}:{timestamp}"
        expected_sig = hmac.new(CLEANUP_SECRET.encode(), payload.encode(), hashlib.sha256).digest()
        actual_sig = base64.urlsafe_b64decode(sig_b64 + "==")
        return hmac.compare_digest(expected_sig, actual_sig)
    except Exception:
        return False


def _is_authenticated_for_files(x_cleanup_token, token_query, session_cookie):
    if not CLEANUP_SECRET:
        return True
    if x_cleanup_token == CLEANUP_SECRET:
        return True
    if token_query == CLEANUP_SECRET:
        return True
    if session_cookie and _verify_session_token(session_cookie):
        return True
    return False


# ============================================================
# SECTION 10 — MCP SERVER & FASTAPI APP
# ============================================================
mcp = FastMCP("sc-report-server", stateless_http=True, json_response=True, host="0.0.0.0")
mcp_asgi_app = mcp.streamable_http_app()


@asynccontextmanager
async def lifespan(app):
    try:
        async with mcp.session_manager.run():
            logger.info("MCP Session started")
            yield
    finally:
        logger.info("App shutting down")


app = FastAPI(
    title="SC Report MCP API & Server",
    description="API & MCP Server สำหรับตรวจสอบสิทธิ์และสร้าง Excel Report จาก BigQuery",
    version="1.4.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["https://ai-uat.scasset.com"],
    allow_credentials=True,
    allow_methods=["GET", "POST", "DELETE"],
    allow_headers=["Authorization", "Content-Type", "current-user", "x-user-email"],
    max_age=3600,
)


# ============================================================
# SECTION 11 — MIDDLEWARE
# ============================================================
class CaptureUsernameMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        username_val = None
        email_source = "none"

        for header_name in ["current-user", "current_user", "x-user-email"]:
            val = request.headers.get(header_name, "").strip()
            if val:
                if is_valid_email(val):
                    username_val = val.lower()
                    email_source = f"header:{header_name}"
                    break
                else:
                    logger.warning(
                        f"[Middleware] Header '{header_name}' has invalid value — skipping"
                    )

        if not username_val:
            auth_header = request.headers.get("authorization", "").strip()
            if auth_header.lower().startswith("bearer "):
                token = auth_header[7:].strip()
                jwt_email = await asyncio.to_thread(extract_email_from_jwt, token)
                if jwt_email:
                    username_val = jwt_email
                    email_source = "jwt"

        masked_user = mask_username(username_val)
        logger.info(
            f"[Middleware] {request.method} {request.url.path} | "
            f"User: {masked_user} | Source: {email_source}"
        )

        token_username = CURRENT_REQUEST_USERNAME.set(username_val)

        bq_username = None
        if username_val:
            bq_username = await lookup_username_async(username_val)
            if bq_username:
                logger.info(f"[Middleware] Resolved BQ username: '{bq_username}' for {masked_user}")
            else:
                logger.warning(f"[Middleware] No BQ username for: {masked_user}")

        token_bq_username = CURRENT_REQUEST_BQ_USERNAME.set(bq_username)
        base_url = str(request.base_url).rstrip("/")
        token_base = CURRENT_REQUEST_BASE_URL.set(base_url)

        try:
            response = await call_next(request)
            return response
        finally:
            CURRENT_REQUEST_USERNAME.reset(token_username)
            CURRENT_REQUEST_BQ_USERNAME.reset(token_bq_username)
            CURRENT_REQUEST_BASE_URL.reset(token_base)


class BlockSSEMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        if request.method == "GET" and request.url.path.rstrip("/") == "/mcp/sse":
            return JSONResponse(
                status_code=405,
                content={"detail": "SSE transport is disabled. Use POST /mcp for Streamable HTTP."},
                headers={"Retry-After": "3600"},
            )
        return await call_next(request)


app.add_middleware(CaptureUsernameMiddleware)
app.add_middleware(BlockSSEMiddleware)


# ============================================================
# SECTION 12 — PYDANTIC MODELS
# ============================================================
AllowedOperators = Literal["=", ">", "<", ">=", "<=", "LIKE", "!="]


class FilterCondition(BaseModel):
    column: str = Field(..., min_length=1, description="ชื่อคอลัมน์ที่จะฟิลเตอร์")
    operator: AllowedOperators = Field(default="=", description="เครื่องหมายเปรียบเทียบ")
    value: str = Field(..., description="ค่าที่ใช้ค้นหา")


class GenerateReportRequest(BaseModel):
    model_config = ConfigDict(extra="allow")
    table_id: str = Field(..., min_length=1)
    limit: int = Field(default=2000, gt=0, le=100000)
    filters: Optional[list[FilterCondition]] = None
    condition: Optional[Literal["AND", "OR"]] = "AND"
    columns: Optional[list[str]] = None
    username: Optional[str] = None
    user_email: Optional[str] = None
    filter_column: Optional[str] = None
    filter_value: Optional[str] = None


class CostEstimateRequest(BaseModel):
    bytes_billed: int = Field(..., ge=0)
    duration_seconds: float = Field(..., gt=0)
    vcpu: float = Field(default=1.0, gt=0)
    memory_gib: float = Field(default=0.5, gt=0)
    usd_to_thb: float = Field(default=32.57, gt=0)


class GenerateChartRequest(BaseModel):
    model_config = ConfigDict(extra="allow")
    table_id: str = Field(..., min_length=1)
    code: str = Field(..., min_length=1)
    limit: int = Field(default=2000, gt=0, le=100000)
    filters: Optional[list[FilterCondition]] = None
    condition: Optional[Literal["AND", "OR"]] = "AND"
    columns: Optional[list[str]] = None
    username: Optional[str] = None
    user_email: Optional[str] = None
    filter_column: Optional[str] = None
    filter_value: Optional[str] = None


# ============================================================
# SECTION 13 — AUTH & RESOLUTION
# ============================================================
def _resolve_username_sync(fallback_email=None):
    bq_username = CURRENT_REQUEST_BQ_USERNAME.get()
    if bq_username and is_valid_user(bq_username):
        return bq_username.strip()

    if fallback_email and is_valid_email(fallback_email):
        db_username = _lookup_username_sync(fallback_email)
        if db_username and is_valid_user(db_username):
            return db_username.strip()

    logger.warning(
        f"[resolve_username] Could not resolve identity — "
        f"middleware_email={repr(CURRENT_REQUEST_USERNAME.get())}, "
        f"bq_username={repr(CURRENT_REQUEST_BQ_USERNAME.get())}, "
        f"fallback_email={repr(fallback_email)}"
    )
    return ""


def resolve_username(fallback_email=None):
    return _resolve_username_sync(fallback_email)


def check_user_permission(username, table_id):
    if not is_valid_user(username):
        return None
    cleaned_table = clean_table_id(table_id)
    if not validate_table_id(cleaned_table):
        return None
    short_name = extract_short_table_name(cleaned_table)
    masked_user = mask_username(username)
    logger.info(f"[Auth] table_id='{table_id}' -> short_name='{short_name}' | user='{masked_user}'")

    if "authenbymenu" in short_name:
        return None

    report_names = TABLE_TO_REPORT_NAMES.get(short_name)
    if not report_names:
        return None

    try:
        clean_username = username.strip()
        placeholders = ", ".join([f"@name{i}" for i in range(len(report_names))])
        query = f"""
            SELECT ReportName
            FROM `{PROJECT_ID}.{DATASET_ID}.{AUTH_TABLE}`
            WHERE UserName = @username
              AND ReportName IN ({placeholders})
            LIMIT 1
        """
        params = [bigquery.ScalarQueryParameter("username", "STRING", clean_username)]
        for i, name in enumerate(report_names):
            params.append(bigquery.ScalarQueryParameter(f"name{i}", "STRING", name.strip()))

        job_config = bigquery.QueryJobConfig(query_parameters=params)
        job = bq_client.query(query, job_config=job_config)
        results = list(job)

        if results:
            return results[0]["ReportName"]
        return None
    except Exception as e:
        logger.error(f"Auth verification database error: {e}")
        return None


def validate_filter_column(table_id, column):
    if not table_id or not column:
        return False
    mapping_dict = get_table_column_mapping(table_id)
    target_column = column.lower()
    return any(k.lower() == target_column for k in mapping_dict.keys())


# ============================================================
# SECTION 14 — AI LOG
# ============================================================
def insert_ai_log(
    username, table_name, report_name, status,
    condition="", size=0.0, bytes_billed=0, url="",
    row_generated=0, ai_log_id=0, cost_thb=0.0, cost_usd=0.0, cost_summary="",
):
    try:
        now = datetime.now(BANGKOK_TZ)
        if ai_log_id == 0:
            ai_log_id = generate_ai_log_id()

        rows_to_insert = [{
            "CreatedAt": now.strftime("%Y-%m-%dT%H:%M:%S"),
            "UserName": (username or "")[:50],
            "TableName": (table_name or "")[:50],
            "ReportName": (report_name or "")[:100],
            "Status": (status or "")[:15],
            "Condition": (condition or "")[:2000],
            "Size": round(size, 4),
            "BytesBilled": bytes_billed,
            "URL": (url or "")[:1000],
            "Row_generated": row_generated,
            "AiLogID": ai_log_id,
            "CostTHB": round(cost_thb, 6),
            "CostUSD": round(cost_usd, 8),
            "CostSummary": (cost_summary or "")[:500],
        }]
        full_table = f"{PROJECT_ID}.{DATASET_ID}.{LOG_TABLE}"
        errors = bq_client.insert_rows_json(full_table, rows_to_insert)
        if errors:
            logger.error(f"AiLog insert FAILED: {errors}")
    except Exception as e:
        logger.error(f"Exception in insert_ai_log: {e}", exc_info=True)


# ============================================================
# SECTION 15 — PARAM MAPPER
# ============================================================
def map_param_type_and_value(column_name, str_val, schema_dict):
    col_type = schema_dict.get(column_name.lower())
    if not col_type:
        return "STRING", str_val

    if col_type in ("INTEGER", "INT64"):
        try:
            val = int(str_val)
            if not (-(2**63) <= val <= 2**63 - 1):
                raise ValueError(f"Value out of INT64 range: {val}")
            return "INT64", val
        except (ValueError, OverflowError):
            return "STRING", str_val

    if col_type in ("FLOAT", "FLOAT64", "NUMERIC", "BIGNUMERIC"):
        try:
            val = float(str_val)
            if not math.isfinite(val):
                raise ValueError(f"Non-finite value: {val}")
            return "FLOAT64", val
        except ValueError:
            return "STRING", str_val

    if col_type in ("BOOLEAN", "BOOL"):
        return "BOOL", str_val.lower() in ("true", "1", "yes")

    return "STRING", str_val


# ============================================================
# SECTION 16 — QUERY BUILDER
# ============================================================
def _validate_columns(short_name, columns):
    if not columns:
        return
    for col in columns:
        if is_sensitive_column(col):
            raise PermissionError(f"ไม่อนุญาตให้ดึงคอลัมน์ที่มีข้อมูลอ่อนไหว: {col}")
        if not validate_filter_column(short_name, col):
            raise PermissionError(f"คอลัมน์ '{col}' ไม่ได้รับอนุญาต")
        sanitize_column_name(col)


def _build_query_and_params(table_id, limit, filter_column, filter_value,
                             filters, columns, condition):
    clean = clean_table_id(table_id)
    if not validate_table_id(clean):
        raise PermissionError(f"โครงสร้าง Table ID '{table_id}' ไม่ถูกต้องหรือไม่อนุญาตให้เข้าถึง")

    short_name = extract_short_table_name(clean)
    if "authenbymenu" in short_name:
        raise PermissionError("Access to the authentication table is strictly prohibited at code level.")

    bq_table_name = TABLE_TO_BQ_TABLE_NAME.get(short_name, clean)
    table_ref = f"{PROJECT_ID}.{DATASET_ID}.{bq_table_name}"
    schema_dict = get_table_schema(table_ref)

    _validate_columns(short_name, columns)

    if columns:
        select_clause = ", ".join([f"`{col}`" for col in columns])
    else:
        select_clause = "*"

    query = f"SELECT {select_clause} FROM `{PROJECT_ID}.{DATASET_ID}.{bq_table_name}`"
    params = []
    where_clauses = []

    if filters:
        for i, f in enumerate(filters):
            if not validate_filter_column(short_name, f.column):
                raise PermissionError(f"คอลัมน์ '{f.column}' ไม่ได้รับอนุญาตสำหรับตารางนี้")
            sanitize_column_name(f.column)
            if f.operator not in ALLOWED_OPERATORS:
                raise PermissionError(f"Operator '{f.operator}' ไม่ได้รับอนุญาต")
            param_name = f"filter_{i}"
            if f.operator == "LIKE":
                where_clauses.append(f"CAST(`{f.column}` AS STRING) LIKE @{param_name}")
                params.append(bigquery.ScalarQueryParameter(param_name, "STRING", f.value))
            else:
                p_type, p_val = map_param_type_and_value(f.column, f.value, schema_dict)
                where_clauses.append(f"`{f.column}` {f.operator} @{param_name}")
                params.append(bigquery.ScalarQueryParameter(param_name, p_type, p_val))
    elif filter_column and filter_value:
        if not validate_filter_column(short_name, filter_column):
            raise PermissionError(f"คอลัมน์ '{filter_column}' ไม่ได้รับอนุญาตสำหรับตารางนี้")
        sanitize_column_name(filter_column)
        p_type, p_val = map_param_type_and_value(filter_column, filter_value, schema_dict)
        where_clauses.append(f"`{filter_column}` = @filter_value")
        params.append(bigquery.ScalarQueryParameter("filter_value", p_type, p_val))

    if where_clauses:
        joiner = " AND "
        if condition and condition.upper() in ("AND", "OR"):
            joiner = f" {condition.upper()} "
        query += " WHERE " + joiner.join(where_clauses)

    safe_limit = min(max(limit, 1), 100000)
    query += f" LIMIT {safe_limit}"

    return query, params, bq_table_name, schema_dict


# ============================================================
# SECTION 17 — FETCH & GENERATE EXCEL  (ใช้ GCS)
# ============================================================
def fetch_and_generate_excel(
    table_id, report_name, limit=2000,
    filter_column=None, filter_value=None,
    filters=None, columns=None, condition="AND",
):
    query, params, bq_table_name, _ = _build_query_and_params(
        table_id, limit, filter_column, filter_value, filters, columns, condition
    )
    actual_sql_executed = query

    job_config = bigquery.QueryJobConfig(query_parameters=params)
    job = bq_client.query(query, job_config=job_config)
    result = job.result()

    bytes_billed = job.total_bytes_billed
    schema = result.schema
    row_count = result.total_rows

    if not row_count:
        return 0, None, bytes_billed, 0.0, actual_sql_executed, generate_ai_log_id()

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Report")

    raw_mapping = get_table_column_mapping(table_id)
    lowercase_mapping = {k.lower(): v for k, v in raw_mapping.items()}

    headers = []
    for field in schema:
        translated_header = lowercase_mapping.get(field.name.lower(), field.name)
        headers.append(translated_header)
    ws.append(headers)

    for row in result:
        row_values = []
        for field in schema:
            val = row[field.name]
            if hasattr(val, "tzinfo") and val.tzinfo is not None:
                val = val.replace(tzinfo=None)
            row_values.append(val)
        ws.append(row_values)

    ai_log_id = generate_ai_log_id()
    secure_file_id = uuid.uuid4().hex[:8]
    timestamp = datetime.now(BANGKOK_TZ).strftime("%Y%m%d_%H%M%S")
    file_name = f"{bq_table_name}_{secure_file_id}_{timestamp}.xlsx"

    # บันทึกลง /tmp ชั่วคราว แล้วอัปโหลด GCS
    local_path = REPORT_DIR / file_name
    wb.save(local_path)
    file_size_mb = round(local_path.stat().st_size / (1024 * 1024), 4)

    upload_to_gcs(local_path, file_name)   # ← อัปโหลด + ลบ /tmp อัตโนมัติ

    return row_count, file_name, bytes_billed, file_size_mb, actual_sql_executed, ai_log_id


# ============================================================
# SECTION 18 — CLEANUP  (ใช้ GCS)
# ============================================================
def cleanup_old_reports():
    deleted = cleanup_gcs_files()
    logger.info(f"Cleanup completed — {deleted} file(s) removed.")


# ============================================================
# SECTION 19 — REPORT GENERATION (orchestrator)
# ============================================================
def _execute_report_generation(
    table_id, limit=2000, filter_column=None, filter_value=None,
    filters=None, columns=None, username=None,
    user_email=None, condition="AND",
):
    _start_time = time.perf_counter()

    tid = clean_table_id(table_id)
    short_name = extract_short_table_name(tid)
    bq_table_name = TABLE_TO_BQ_TABLE_NAME.get(short_name, tid)

    username_to_use = _resolve_username_sync(fallback_email=user_email)

    if not username_to_use:
        logger.warning(
            f"[Auth] REJECTED | middleware_email={repr(CURRENT_REQUEST_USERNAME.get())} | "
            f"bq_username={repr(CURRENT_REQUEST_BQ_USERNAME.get())} | ai_username={repr(username)}"
        )
        insert_ai_log("UNAUTH", bq_table_name, "", "FAIL_UNAUTH")
        return {"success": False, "message": "ไม่สามารถยืนยันตัวตนได้ กรุณา Login ใหม่อีกครั้ง"}

    if "authenbymenu" in short_name:
        insert_ai_log(username_to_use, bq_table_name, "", "FAIL")
        return {"success": False, "message": "[Access Denied] ระบบไม่อนุญาตให้เข้าถึงตารางสิทธิ์"}

    if not validate_table_id(tid):
        insert_ai_log(username_to_use, bq_table_name, "", "FAIL")
        return {"success": False, "message": f"table_id '{tid}' ไม่ถูกต้อง"}

    report_name = check_user_permission(username_to_use, tid)
    if not report_name:
        insert_ai_log(username_to_use, bq_table_name, "", "FAIL")
        return {"success": False, "message": "ขออภัย คุณไม่มีสิทธิ์เข้าถึงรายงานนี้"}

    try:
        row_count, file_name, bytes_billed, file_size_mb, actual_sql, ai_log_id = (
            fetch_and_generate_excel(tid, report_name, limit,
                                     filter_column=filter_column, filter_value=filter_value,
                                     filters=filters, columns=columns, condition=condition)
        )

        _elapsed = time.perf_counter() - _start_time
        cost = calculate_total_cost(
            bytes_billed=bytes_billed, duration_seconds=_elapsed,
            vcpu=COST_CLOUD_RUN_VCPU, memory_gib=COST_CLOUD_RUN_MEM_GIB,
            usd_to_thb=COST_USD_TO_THB, bq_price_per_tb_usd=COST_BQ_PRICE_PER_TB,
        )
        cost_breakdown = format_cost_breakdown(cost)

        if not row_count:
            insert_ai_log(username_to_use, bq_table_name, report_name, "FAIL",
                          condition=actual_sql, bytes_billed=bytes_billed, ai_log_id=ai_log_id,
                          cost_thb=cost.grand_total_thb, cost_usd=cost.grand_total_usd,
                          cost_summary=cost.summary)
            return {"success": False, "message": f"ไม่พบข้อมูลในรายงาน {report_name}"}

        base_url = CURRENT_REQUEST_BASE_URL.get() or PUBLIC_BASE_URL
        download_url = f"{base_url.rstrip('/')}/download/{file_name}"

        insert_ai_log(username=username_to_use, table_name=bq_table_name, report_name=report_name,
                      status="OK", condition=actual_sql, size=file_size_mb,
                      bytes_billed=bytes_billed, url=download_url,
                      row_generated=row_count or 0, ai_log_id=ai_log_id,
                      cost_thb=cost.grand_total_thb, cost_usd=cost.grand_total_usd,
                      cost_summary=cost.summary)
        return {
            "success": True,
            "message": f"จัดเตรียมรายงาน{report_name} ({row_count} แถว) สำเร็จ",
            "download_url": download_url, "file_name": file_name,
            "row_count": row_count, "file_size_mb": file_size_mb,
            "report_name": report_name, "cost": cost_breakdown,
        }

    except PermissionError as pe:
        insert_ai_log(username_to_use, bq_table_name, report_name or "", "FAIL")
        return {"success": False, "message": str(pe)}
    except Exception as e:
        logger.error(f"Execution Error: {e}", exc_info=True)
        insert_ai_log(username_to_use, bq_table_name, report_name or "", "FAIL")
        return {"success": False, "message": "พบปัญหาการสร้าง Excel"}


def fetch_dataframe_for_chart(
    table_id, limit=2000, filter_column=None, filter_value=None,
    filters=None, columns=None, condition="AND",
):
    query, params, bq_table_name, _ = _build_query_and_params(
        table_id, limit, filter_column, filter_value, filters, columns, condition
    )
    actual_sql_executed = query
    job_config = bigquery.QueryJobConfig(query_parameters=params)
    job = bq_client.query(query, job_config=job_config)
    df = job.to_dataframe()
    bytes_billed = job.total_bytes_billed
    return df, bytes_billed, actual_sql_executed


def _execute_chart_generation(
    table_id, code, limit=2000, filter_column=None, filter_value=None,
    filters=None, columns=None, username=None,
    user_email=None, condition="AND",
) -> dict:
    _start_time = time.perf_counter()

    tid = clean_table_id(table_id)
    short_name = extract_short_table_name(tid)
    bq_table_name = TABLE_TO_BQ_TABLE_NAME.get(short_name, tid)

    username_to_use = _resolve_username_sync(fallback_email=user_email)

    if not username_to_use:
        logger.warning(
            f"[Auth] REJECTED | middleware_email={repr(CURRENT_REQUEST_USERNAME.get())} | "
            f"bq_username={repr(CURRENT_REQUEST_BQ_USERNAME.get())} | ai_username={repr(username)}"
        )
        insert_ai_log("UNAUTH", bq_table_name, "", "FAIL_UNAUTH")
        return {"success": False, "message": "ไม่สามารถยืนยันตัวตนได้ กรุณา Login ใหม่อีกครั้ง"}

    if "authenbymenu" in short_name:
        insert_ai_log(username_to_use, bq_table_name, "", "FAIL_CHART")
        return {"success": False, "message": "[Access Denied] ระบบไม่อนุญาตให้เข้าถึงตารางสิทธิ์"}

    if not validate_table_id(tid):
        insert_ai_log(username_to_use, bq_table_name, "", "FAIL_CHART")
        return {"success": False, "message": f"table_id '{tid}' ไม่ถูกต้อง"}

    report_name = check_user_permission(username_to_use, tid)
    if not report_name:
        insert_ai_log(username_to_use, bq_table_name, "", "FAIL_CHART")
        return {"success": False, "message": "ขออภัย คุณไม่มีสิทธิ์เข้าถึงรายงานนี้"}

    try:
        df, bytes_billed, actual_sql = fetch_dataframe_for_chart(
            tid, limit, filter_column, filter_value, filters, columns, condition
        )

        if df.empty:
            ai_log_id = generate_ai_log_id()
            _elapsed = time.perf_counter() - _start_time
            cost = calculate_total_cost(bytes_billed, _elapsed,
                                         COST_CLOUD_RUN_VCPU, COST_CLOUD_RUN_MEM_GIB,
                                         COST_USD_TO_THB, COST_BQ_PRICE_PER_TB)
            insert_ai_log(username_to_use, bq_table_name, report_name, "FAIL_CHART",
                          condition=actual_sql, bytes_billed=bytes_billed, ai_log_id=ai_log_id,
                          cost_thb=cost.grand_total_thb, cost_usd=cost.grand_total_usd,
                          cost_summary=cost.summary)
            return {"success": False, "message": f"ไม่พบข้อมูลในรายงาน {report_name} เพื่อใช้วาดกราฟ"}

        raw_mapping = get_table_column_mapping(tid)
        lowercase_mapping = {k.lower(): v for k, v in raw_mapping.items()}
        rename_dict = {}
        for col in df.columns:
            translated = lowercase_mapping.get(col.lower(), col)
            rename_dict[col] = translated
        df = df.rename(columns=rename_dict)

        ai_log_id = generate_ai_log_id()
        secure_file_id = uuid.uuid4().hex[:8]
        timestamp = datetime.now(BANGKOK_TZ).strftime("%Y%m%d_%H%M%S")
        file_name = f"chart_{bq_table_name}_{secure_file_id}_{timestamp}.png"
        local_path = REPORT_DIR / file_name

        res = run_visualization_code(df, code, str(local_path))

        _elapsed = time.perf_counter() - _start_time
        cost = calculate_total_cost(bytes_billed=bytes_billed, duration_seconds=_elapsed,
                                     vcpu=COST_CLOUD_RUN_VCPU, memory_gib=COST_CLOUD_RUN_MEM_GIB,
                                     usd_to_thb=COST_USD_TO_THB, bq_price_per_tb_usd=COST_BQ_PRICE_PER_TB)
        cost_breakdown = format_cost_breakdown(cost)

        if not res["success"]:
            insert_ai_log(username_to_use, bq_table_name, report_name, "FAIL_CHART",
                          condition=actual_sql, bytes_billed=bytes_billed, ai_log_id=ai_log_id,
                          cost_thb=cost.grand_total_thb, cost_usd=cost.grand_total_usd,
                          cost_summary=cost.summary)
            sample_data = df.head(3).to_dict('records')
            return {
                "success": False,
                "message": f"วาดกราฟไม่สำเร็จ: {res['message']}",
                "available_columns": list(df.columns),
                "sample_data": sample_data,
                "row_count": len(df),
                "hint": "ใช้ชื่อ column จาก available_columns เท่านั้น (ชื่อภาษาไทยที่แปลแล้ว)"
            }

        file_size_mb = round(local_path.stat().st_size / (1024 * 1024), 4)

        upload_to_gcs(local_path, file_name)   # ← อัปโหลด + ลบ /tmp อัตโนมัติ

        base_url = CURRENT_REQUEST_BASE_URL.get() or PUBLIC_BASE_URL
        download_url = f"{base_url.rstrip('/')}/download/{file_name}"

        insert_ai_log(username=username_to_use, table_name=bq_table_name, report_name=report_name,
                      status="OK_CHART", condition=actual_sql, size=file_size_mb,
                      bytes_billed=bytes_billed, url=download_url,
                      row_generated=len(df), ai_log_id=ai_log_id,
                      cost_thb=cost.grand_total_thb, cost_usd=cost.grand_total_usd,
                      cost_summary=cost.summary)

        return {
            "success": True,
            "message": f"จัดเตรียมรายงานแผนภูมิของ{report_name} สำเร็จ",
            "download_url": download_url, "file_name": file_name,
            "row_count": len(df), "file_size_mb": file_size_mb,
            "report_name": report_name, "cost": cost_breakdown,
            "available_columns": list(df.columns),
        }

    except Exception as e:
        logger.error(f"Chart Execution Error: {e}", exc_info=True)
        insert_ai_log(username_to_use, bq_table_name, report_name or "", "FAIL_CHART")
        return {"success": False, "message": f"พบปัญหาการดึงข้อมูลเพื่อวาดกราฟ: {str(e)}"}


# ============================================================
# SECTION 20 — MCP TOOLS
# ============================================================
@mcp.tool(name="sc_report_export", description="ดึงข้อมูล SC Report และสร้างไฟล์ Excel")
def mcp_sc_report_export(
    table_id: str,
    username: str,
    limit: int = 2000,
    filter_column: Optional[str] = None,
    filter_value: Optional[str] = None,
    filters: Optional[list[dict]] = None,
    columns: Optional[list[str]] = None,
    **kwargs,
) -> dict:
    parsed_filters = None
    if filters:
        try:
            parsed_filters = [FilterCondition(**f) for f in filters]
        except Exception:
            return {"success": False, "message": "รูปแบบ filters ไม่ถูกต้อง"}

    res = _execute_report_generation(
        table_id=table_id, limit=limit, filter_column=filter_column,
        filter_value=filter_value, filters=parsed_filters, columns=columns,
        username=username, condition="AND",
    )
    if res.get("success") and "cost" in res:
        res["cost_summary"] = res["cost"].get("summary", "")
    return res


@mcp.tool(name="generate_excel_report", description="สร้างไฟล์ Excel จากรายงานใน BigQuery")
def mcp_generate_excel_report(
    table_id: str,
    limit: int = 2000,
    filter_column: Optional[str] = None,
    filter_value: Optional[str] = None,
    filters: Optional[list[dict]] = None,
    condition: Optional[str] = None,
    username: Optional[str] = None,
    user_email: Optional[str] = None,
    columns: Optional[list[str]] = None,
    **kwargs,
) -> str:
    parsed_filters = None
    if filters:
        try:
            parsed_filters = [FilterCondition(**f) for f in filters]
        except Exception:
            return "รูปแบบ filters ไม่ถูกต้อง"

    res = _execute_report_generation(
        table_id=table_id, limit=limit, filter_column=filter_column,
        filter_value=filter_value, filters=parsed_filters, columns=columns,
        username=username, user_email=user_email, condition=condition or "AND",
    )
    if res["success"]:
        cost_line = ""
        if "cost" in res:
            c = res["cost"]
            cost_line = (
                f"ค่าใช้จ่าย: BQ ${c['bigquery']['cost_usd']:.6f} | "
                f"Cloud Run ${c['cloud_run']['total_cost_usd']:.6f} | "
                f"รวม ฿{c['grand_total_thb']:.4f}"
            )
        return (
            f"จัดเตรียม **รายงาน{res['report_name']}** ({res['row_count']} แถว) เรียบร้อยแล้วครับ\n"
            f"{res['download_url']}\n{cost_line}"
        )
    return res["message"]


@mcp.tool(
    name="generate_chart_report",
    description=(
        "สร้างกราฟหรือแผนภูมิจากข้อมูลรายงานใน BigQuery ด้วยโค้ด Python (Seaborn/Matplotlib)\n\n"
        "⚠️ สำคัญ: \n"
        "- ชื่อ column ใน DataFrame จะถูกแปลเป็นภาษาไทยตาม data dictionary อัตโนมัติ\n"
        "- ต้องใช้ชื่อ column ภาษาไทยในโค้ด เช่น 'ชื่อกลุ่มงาน' ไม่ใช่ 'WorkGroupName'\n"
        "- ถ้าโค้ด error ระบบจะ return available_columns ให้ใช้ชื่อจากตรงนั้น\n"
        "- โค้ดต้องมี plt.savefig() หรือ savefig() ก่อน plt.show()"
    )
)
def mcp_generate_chart_report(
    table_id: str,
    code: str,
    limit: int = 2000,
    filter_column: Optional[str] = None,
    filter_value: Optional[str] = None,
    filters: Optional[list[dict]] = None,
    condition: Optional[str] = None,
    username: Optional[str] = None,
    user_email: Optional[str] = None,
    columns: Optional[list[str]] = None,
    **kwargs,
) -> str:
    parsed_filters = None
    if filters:
        try:
            parsed_filters = [FilterCondition(**f) for f in filters]
        except Exception:
            return "รูปแบบ filters ไม่ถูกต้อง"

    res = _execute_chart_generation(
        table_id=table_id, code=code, limit=limit, filter_column=filter_column,
        filter_value=filter_value, filters=parsed_filters, columns=columns,
        username=username, user_email=user_email, condition=condition or "AND",
    )

    if res["success"]:
        cost_line = ""
        if "cost" in res:
            c = res["cost"]
            cost_line = (
                f"ค่าใช้จ่าย: BQ ${c['bigquery']['cost_usd']:.6f} | "
                f"Cloud Run ${c['cloud_run']['total_cost_usd']:.6f} | "
                f"รวม ฿{c['grand_total_thb']:.4f}"
            )
        cols_hint = ""
        if "available_columns" in res:
            cols_hint = f"\n\n📊 Columns ที่ใช้ได้: {res['available_columns']}"

        return (
            f"จัดเตรียม **แผนภูมิของรายงาน{res['report_name']}** ({res['row_count']} แถว) เรียบร้อยแล้วครับ\n"
            f"{res['download_url']}\n{cost_line}{cols_hint}"
        )

    if "available_columns" in res:
        sample_str = ""
        if res.get("sample_data"):
            sample_str = f"\n\n📋 ตัวอย่างข้อมูล 3 แถวแรก:\n```\n{res['sample_data']}\n```"
        return (
            f"❌ {res['message']}\n\n"
            f"📊 **Columns ที่ใช้ได้ใน DataFrame** (ชื่อภาษาไทย):\n"
            f"```\n{res['available_columns']}\n```"
            f"{sample_str}\n\n"
            f"💡 กรุณาใช้ชื่อ column จาก available_columns เท่านั้น"
        )

    return res["message"]


@mcp.tool(
    name="check_accessible_reports",
    description="ตรวจสอบว่าผู้ใช้งานคนปัจจุบันมีสิทธิ์เข้าถึงรายงาน (ReportName / Table) ใดบ้าง ไม่ต้องใส่พารามิเตอร์ใดๆ ระบบจะดึงจากอีเมลผู้ใช้อัตโนมัติจาก Header"
)
def mcp_check_accessible_reports() -> dict:
    username = resolve_username(None)
    if not username or username == "Anonymous" or username == "Unknown":
        return {"success": False, "message": "ไม่พบข้อมูลสิทธิ์ของผู้ใช้งาน หรือ ไม่สามารถระบุตัวตนผู้ใช้ได้"}

    try:
        query = f"""
            SELECT ReportName 
            FROM `{PROJECT_ID}.{DATASET_ID}.{AUTH_TABLE}`
            WHERE UserName = @username
        """
        job_config = bigquery.QueryJobConfig(
            query_parameters=[
                bigquery.ScalarQueryParameter("username", "STRING", username)
            ]
        )
        results = list(bq_client.query(query, job_config=job_config).result())
        reports = [row["ReportName"] for row in results]

        return {
            "success": True,
            "username": username,
            "accessible_reports": reports,
            "total": len(reports)
        }
    except Exception as e:
        logger.error(f"Error fetching accessible reports: {e}")
        return {"success": False, "message": f"เกิดข้อผิดพลาดในการดึงข้อมูลสิทธิ์: {str(e)}"}


# ============================================================
# SECTION 21 — REST API ENDPOINTS
# ============================================================
@app.post("/generate_excel_report")
def generate_excel_report(request: GenerateReportRequest):
    res = _execute_report_generation(
        table_id=request.table_id, limit=request.limit,
        filter_column=request.filter_column, filter_value=request.filter_value,
        filters=request.filters, columns=request.columns,
        username=request.username, user_email=request.user_email,
        condition=request.condition or "AND",
    )
    if res["success"]:
        return {
            "success": True,
            "message": f"จัดเตรียม **รายงาน{res['report_name']}** สำเร็จ\n{res['download_url']}",
            "cost": res.get("cost"),
        }
    return {"success": False, "message": res["message"]}


@app.post("/generate_chart_report")
def generate_chart_report(request: GenerateChartRequest):
    res = _execute_chart_generation(
        table_id=request.table_id, code=request.code, limit=request.limit,
        filter_column=request.filter_column, filter_value=request.filter_value,
        filters=request.filters, columns=request.columns,
        username=request.username, user_email=request.user_email,
        condition=request.condition or "AND",
    )
    if res["success"]:
        return {
            "success": True,
            "message": f"จัดเตรียม **แผนภูมิของ{res['report_name']}** สำเร็จ\n{res['download_url']}",
            "download_url": res["download_url"],
            "cost": res.get("cost"),
        }
    return {"success": False, "message": res["message"]}


@app.post("/cost_estimate")
def cost_estimate(req: CostEstimateRequest):
    cost = calculate_total_cost(
        bytes_billed=req.bytes_billed, duration_seconds=req.duration_seconds,
        vcpu=req.vcpu, memory_gib=req.memory_gib, usd_to_thb=req.usd_to_thb,
        bq_price_per_tb_usd=COST_BQ_PRICE_PER_TB,
    )
    return format_cost_breakdown(cost)


@app.get("/download/{file_name}")
async def download_file(
    file_name: str,
    direct: bool = False,
    x_cleanup_token: Optional[str] = Header(None),
    token: Optional[str] = None,
    sc_report_session: Optional[str] = Cookie(None),
):
    if not _is_authenticated_for_files(x_cleanup_token, token, sc_report_session):
        raise HTTPException(status_code=401, detail="Unauthorized")

    safe_name = _safe_report_filename(file_name)

    if not direct:
        return JSONResponse({
            "success": True,
            "file_name": safe_name,
            "download_url": generate_gcs_signed_url(safe_name),
            "expires_in_seconds": GCS_SIGNED_URL_TTL,
        })

    blob = get_gcs_blob(safe_name)
    stream = blob.open("rb")

    disposition = "inline" if safe_name.endswith(".png") else "attachment"

    headers = {
        "Content-Disposition": f'{disposition}; filename="{quote(safe_name)}"',
        "X-Content-Type-Options": "nosniff",
        "Cache-Control": "private, max-age=300",
    }

    return StreamingResponse(
        stream,
        media_type=get_gcs_content_type(safe_name),
        headers=headers,
    )


@app.get("/health")
def health_check():
    return {"status": "healthy", "storage": "gcs", "bucket": REPORT_BUCKET}


# ============================================================
# SECTION 22 — FILE MANAGER  (ใช้ GCS)
# ============================================================
@app.post("/files/login")
async def files_login(
    x_cleanup_token: Optional[str] = Header(None),
    token: Optional[str] = None,
):
    if not CLEANUP_SECRET:
        response = JSONResponse({"success": True, "message": "Logged in (no auth required)"})
        response.set_cookie(
            key=SESSION_COOKIE_NAME, value="no-auth",
            max_age=SESSION_MAX_AGE, httponly=True, secure=False, samesite="strict",
        )
        return response

    if (x_cleanup_token or token) != CLEANUP_SECRET:
        raise HTTPException(status_code=401, detail="Invalid token")

    session = _make_session_token()
    response = JSONResponse({"success": True, "message": "Logged in"})
    response.set_cookie(
        key=SESSION_COOKIE_NAME, value=session,
        max_age=SESSION_MAX_AGE, httponly=True, secure=True, samesite="strict",
    )
    return response


@app.post("/files/logout")
async def files_logout():
    response = JSONResponse({"success": True, "message": "Logged out"})
    response.delete_cookie(SESSION_COOKIE_NAME)
    return response


def _render_login_page():
    return """<!DOCTYPE html>
<html><head><meta charset="UTF-8"><title>Login - SC Report</title>
<style>
  body { background: #0d1117; color: #f0f6fc; font-family: -apple-system, sans-serif;
         display: flex; justify-content: center; align-items: center; height: 100vh; margin: 0; }
  .box { background: rgba(22,27,34,0.9); border: 1px solid #30363d; border-radius: 12px;
         padding: 2rem; max-width: 400px; width: 100%; }
  h1 { margin-top: 0; }
  input { width: 100%; padding: 0.75rem; background: #0d1117; border: 1px solid #30363d;
          color: #f0f6fc; border-radius: 8px; font-size: 1rem; box-sizing: border-box; }
  button { margin-top: 1rem; width: 100%; padding: 0.75rem; background: #58a6ff;
           color: #0d1117; border: none; border-radius: 8px; font-weight: 600;
           font-size: 1rem; cursor: pointer; }
  button:hover { opacity: 0.9; }
  .error { color: #ff7b72; margin-top: 0.5rem; display: none; }
</style></head>
<body><div class="box">
  <h1>🔐 SC Report Login</h1>
  <p style="color: #8b949e;">กรอก cleanup token เพื่อเข้าถึง file manager</p>
  <input type="password" id="tokenInput" placeholder="Token" autofocus>
  <button onclick="doLogin()">Login</button>
  <p class="error" id="errorMsg">Token ไม่ถูกต้อง</p>
</div>
<script>
async function doLogin() {
  const token = document.getElementById('tokenInput').value;
  const res = await fetch('/files/login', {
    method: 'POST',
    headers: { 'x-cleanup-token': token }
  });
  if (res.ok) {
    window.location.href = '/files';
  } else {
    document.getElementById('errorMsg').style.display = 'block';
  }
}
document.getElementById('tokenInput').addEventListener('keypress', (e) => {
  if (e.key === 'Enter') doLogin();
});
</script></body></html>"""


@app.get("/files")
@app.get("/file")
def list_files(
    request: Request,
    x_cleanup_token: Optional[str] = Header(None),
    token: Optional[str] = None,
    format: Optional[str] = None,
    sc_report_session: Optional[str] = Cookie(None),
):
    if not _is_authenticated_for_files(x_cleanup_token, token, sc_report_session):
        accept = request.headers.get("accept", "")
        if "text/html" in accept:
            return HTMLResponse(content=_render_login_page(), status_code=401)
        raise HTTPException(status_code=401, detail="Unauthorized")

    try:
        files = list_gcs_files()   # ← GCS
        total_size_bytes = sum(f["size_bytes"] for f in files)

        if format == "json" or "text/html" not in request.headers.get("accept", ""):
            return {
                "bucket": REPORT_BUCKET,
                "prefix": GCS_REPORT_PREFIX,
                "files": files,
            }

        html_template = f"""<!DOCTYPE html>
<html lang="th">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>SC Report - File Manager</title>
    <link href="https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;500;600;700&family=Sarabun:wght@300;400;500;600;700&display=swap" rel="stylesheet">
    <style>
        :root {{
            --bg-color: #0d1117; --card-bg: rgba(22,27,34,0.7); --border-color: rgba(48,54,61,0.8);
            --text-main: #f0f6fc; --text-muted: #8b949e; --accent-blue: #58a6ff;
            --accent-green: #3fb950; --accent-red: #ff7b72; --glass-blur: blur(12px);
        }}
        * {{ box-sizing: border-box; margin: 0; padding: 0; }}
        body {{ background-color: var(--bg-color); color: var(--text-main);
                font-family: 'Outfit','Sarabun',sans-serif; min-height: 100vh;
                background-image: radial-gradient(circle at 10% 20%, rgba(90,120,250,0.05) 0%, transparent 40%),
                                  radial-gradient(circle at 90% 80%, rgba(250,100,100,0.03) 0%, transparent 40%);
                background-attachment: fixed; padding: 2rem 1.5rem; }}
        .container {{ max-width: 1200px; margin: 0 auto; }}
        header {{ display: flex; justify-content: space-between; align-items: center;
                  margin-bottom: 2rem; flex-wrap: wrap; gap: 1rem; }}
        .title-area h1 {{ font-size: 2rem; font-weight: 700;
                          background: linear-gradient(135deg, #fff 0%, #8b949e 100%);
                          -webkit-background-clip: text; -webkit-text-fill-color: transparent; }}
        .title-area p {{ color: var(--text-muted); margin-top: 0.25rem; font-size: 0.9rem; }}
        .stats-bar {{ display: flex; gap: 1.5rem; }}
        .stat-card {{ background: var(--card-bg); border: 1px solid var(--border-color);
                      backdrop-filter: var(--glass-blur); border-radius: 12px; padding: 0.75rem 1.25rem;
                      display: flex; flex-direction: column; }}
        .stat-card .value {{ font-size: 1.4rem; font-weight: 600; color: var(--accent-blue); }}
        .stat-card .label {{ font-size: 0.75rem; text-transform: uppercase;
                             letter-spacing: 0.05em; color: var(--text-muted); margin-top: 0.15rem; }}
        .toolbar {{ background: var(--card-bg); border: 1px solid var(--border-color);
                    backdrop-filter: var(--glass-blur); border-radius: 16px; padding: 1rem;
                    margin-bottom: 1.5rem; display: flex; justify-content: space-between;
                    align-items: center; flex-wrap: wrap; gap: 1rem; }}
        .search-wrapper {{ position: relative; flex: 1; max-width: 400px; }}
        .search-input {{ width: 100%; background: rgba(13,17,23,0.8); border: 1px solid var(--border-color);
                         border-radius: 8px; padding: 0.6rem 1rem 0.6rem 2.5rem; color: var(--text-main);
                         font-family: inherit; font-size: 0.95rem; transition: all 0.2s; }}
        .search-input:focus {{ outline: none; border-color: var(--accent-blue); box-shadow: 0 0 0 3px rgba(88,166,255,0.15); }}
        .search-icon {{ position: absolute; left: 0.8rem; top: 50%; transform: translateY(-50%);
                        fill: var(--text-muted); width: 18px; height: 18px; pointer-events: none; }}
        .btn {{ background: var(--accent-blue); color: #0d1117; border: none; border-radius: 8px;
                padding: 0.6rem 1.2rem; font-weight: 600; font-size: 0.9rem; cursor: pointer;
                transition: all 0.2s; display: inline-flex; align-items: center; gap: 0.5rem; font-family: inherit; }}
        .btn:hover {{ opacity: 0.9; transform: translateY(-1px); }}
        .btn-outline {{ background: transparent; border: 1px solid var(--border-color); color: var(--text-main); }}
        .btn-outline:hover {{ background: rgba(255,255,255,0.05); }}
        .btn-danger {{ background: rgba(255,123,114,0.15); border: 1px solid rgba(255,123,114,0.4);
                       color: var(--accent-red); }}
        .btn-danger:hover {{ background: var(--accent-red); color: #0d1117; border-color: var(--accent-red); }}
        .file-grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(280px,1fr)); gap: 1.25rem; }}
        .file-card {{ background: var(--card-bg); border: 1px solid var(--border-color); border-radius: 16px;
                      padding: 1.25rem; display: flex; flex-direction: column; justify-content: space-between;
                      transition: all 0.25s cubic-bezier(0.4,0,0.2,1); }}
        .file-card:hover {{ transform: translateY(-4px); border-color: var(--accent-blue);
                            box-shadow: 0 8px 24px rgba(0,0,0,0.3); }}
        .file-header {{ display: flex; align-items: flex-start; gap: 0.75rem; margin-bottom: 1rem; }}
        .file-icon {{ width: 44px; height: 44px; border-radius: 10px; display: flex;
                      align-items: center; justify-content: center; flex-shrink: 0;
                      border: 1px solid rgba(255,255,255,0.05); }}
        .icon-xlsx {{ color: var(--accent-green); background: rgba(63,185,80,0.08);
                      border-color: rgba(63,185,80,0.15); }}
        .icon-png {{ color: var(--accent-blue); background: rgba(88,166,255,0.08);
                     border-color: rgba(88,166,255,0.15); }}
        .file-info {{ min-width: 0; }}
        .file-name {{ font-size: 0.95rem; font-weight: 600; margin-bottom: 0.25rem;
                      white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }}
        .file-meta {{ font-size: 0.75rem; color: var(--text-muted); display: flex;
                      flex-direction: column; gap: 0.15rem; }}
        .thumbnail-container {{ width: 100%; height: 140px; border-radius: 10px; background: #07090e;
                                 overflow: hidden; margin-bottom: 1rem;
                                 border: 1px solid rgba(255,255,255,0.05);
                                 display: flex; align-items: center; justify-content: center; cursor: zoom-in; }}
        .thumbnail-img {{ max-width: 100%; max-height: 100%; object-fit: contain; transition: transform 0.2s; }}
        .thumbnail-container:hover .thumbnail-img {{ transform: scale(1.05); }}
        .file-actions {{ display: flex; gap: 0.5rem; margin-top: auto; }}
        .action-btn {{ flex: 1; padding: 0.5rem; border-radius: 8px; font-size: 0.85rem; font-weight: 500;
                       text-decoration: none; text-align: center; transition: all 0.2s; cursor: pointer;
                       border: 1px solid transparent; }}
        .btn-view {{ background: rgba(88,166,255,0.1); color: var(--accent-blue); border-color: rgba(88,166,255,0.2); }}
        .btn-view:hover {{ background: var(--accent-blue); color: #0d1117; }}
        .btn-download {{ background: var(--accent-blue); color: #0d1117; font-weight: 600; }}
        .btn-download:hover {{ opacity: 0.9; }}
        .btn-delete {{ background: rgba(255,123,114,0.1); color: var(--accent-red); border-color: rgba(255,123,114,0.2); }}
        .btn-delete:hover {{ background: var(--accent-red); color: #0d1117; }}
        .empty-state {{ background: var(--card-bg); border: 2px dashed var(--border-color); border-radius: 20px;
                        padding: 4rem 2rem; text-align: center; color: var(--text-muted); grid-column: 1 / -1; }}
        .empty-state h3 {{ color: var(--text-main); font-size: 1.3rem; margin-bottom: 0.5rem; }}
        .modal {{ display: none; position: fixed; z-index: 1000; left: 0; top: 0; width: 100%; height: 100%;
                  background: rgba(13,17,23,0.9); backdrop-filter: blur(8px);
                  align-items: center; justify-content: center; }}
        .modal-content {{ max-width: 90%; max-height: 85%; box-shadow: 0 24px 48px rgba(0,0,0,0.5);
                          border-radius: 12px; border: 1px solid var(--border-color); background: #0d1117;
                          overflow: hidden; display: flex; flex-direction: column; }}
        .modal-body {{ padding: 1rem; display: flex; justify-content: center; align-items: center; }}
        .modal-body img {{ max-width: 100%; max-height: 70vh; object-fit: contain; }}
        .modal-header {{ display: flex; justify-content: space-between; align-items: center;
                         padding: 1rem 1.5rem; border-bottom: 1px solid var(--border-color); }}
        .modal-title {{ font-size: 1.1rem; font-weight: 600; }}
        .close-btn {{ color: var(--text-muted); font-size: 28px; font-weight: bold; cursor: pointer; }}
        .close-btn:hover {{ color: var(--text-main); }}
        .gcs-badge {{ background: rgba(88,166,255,0.1); border: 1px solid rgba(88,166,255,0.25);
                      color: var(--accent-blue); border-radius: 6px; padding: 0.2rem 0.6rem;
                      font-size: 0.75rem; font-weight: 600; }}
    </style>
</head>
<body>
<div class="container">
    <header>
        <div class="title-area">
            <h1>📁 File Manager <span class="gcs-badge">☁️ GCS</span></h1>
            <p>gs://{REPORT_BUCKET}/{GCS_REPORT_PREFIX} — จัดการไฟล์รายงานและกราฟแดชบอร์ด</p>
        </div>
        <div class="stats-bar">
            <div class="stat-card">
                <span class="value">{len(files)}</span>
                <span class="label">จำนวนไฟล์ทั้งหมด</span>
            </div>
            <div class="stat-card">
                <span class="value">{round(total_size_bytes / (1024*1024), 2)} MB</span>
                <span class="label">ขนาดรวมทั้งหมด</span>
            </div>
        </div>
    </header>
    <div class="toolbar">
        <div class="search-wrapper">
            <svg class="search-icon" viewBox="0 0 16 16"><path d="M10.68 11.74a6 6 0 0 1-8.322-8.322 6 6 0 0 1 8.322 8.322Zm1.06-1.06A7.5 7.5 0 1 0 2.22 2.22a7.5 7.5 0 0 0 10.56 10.56L15 15.28a.749.749 0 0 0 1.06-1.06l-4.32-4.32Z"></path></svg>
            <input type="text" id="searchInput" class="search-input" placeholder="ค้นหาชื่อไฟล์..." onkeyup="filterFiles()">
        </div>
        <div style="display:flex;gap:0.5rem;">
            <button class="btn btn-outline" onclick="window.location.reload()">🔄 รีเฟรช</button>
            <button class="btn btn-danger" onclick="triggerCleanup()">🧹 เคลียร์ไฟล์ขยะ</button>
        </div>
    </div>
    <div class="file-grid" id="fileGrid">
"""

        if not files:
            html_template += """
        <div class="empty-state">
            <h3>ไม่พบไฟล์ใน GCS Bucket</h3>
            <p>ยังไม่มีไฟล์รายงานหรือกราฟบันทึกไว้ใน bucket นี้</p>
        </div>
"""
        else:
            for f in files:
                name = _safe_report_filename(f["name"])
                name_html = html.escape(name, quote=True)
                name_attr = html.escape(name.lower(), quote=True)
                name_js = json.dumps(name)

                download_path = f"/download/{quote(name)}?direct=true"
                download_path_html = html.escape(download_path, quote=True)

                size_mb = f["size_mb"]
                raw_mod = f["modified_at"]
                try:
                    mod_display = datetime.fromisoformat(raw_mod).strftime("%Y-%m-%d %H:%M:%S")
                except Exception:
                    mod_display = raw_mod
                is_img = name.endswith(".png")
                icon_cls = "icon-png" if is_img else "icon-xlsx"
                icon_svg = (
                    '<svg viewBox="0 0 16 16" width="20" height="20" fill="currentColor">'
                    '<path d="M2.002 1a2 2 0 0 0-2 2v10a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V3a2 2 0 0 0-2-2h-12zm12 1a1 1 0 0 1 1 1v6.5l-3.777-1.947a.5.5 0 0 0-.577.093l-3.71 3.71-2.66-1.772a.5.5 0 0 0-.63.062L1.002 12V3a1 1 0 0 1 1-1h12z"/>'
                    '<path d="M10.5 7a1.5 1.5 0 1 0 0-3 1.5 1.5 0 0 0 0 3z"/></svg>'
                    if is_img else
                    '<svg viewBox="0 0 16 16" width="20" height="20" fill="currentColor">'
                    '<path d="M14 14V4.5L9.5 0H4a2 2 0 0 0-2 2v12a2 2 0 0 0 2 2h8a2 2 0 0 0 2-2zM9.5 1.5 13 5h-3.5V1.5zM3 9h10v1H3V9zm0 2h10v1H3v-1zm0 2h7v1H3v-1z"/></svg>'
                )
                html_template += f"""
        <div class="file-card" data-name="{name_attr}">
            <div class="file-header">
                <div class="file-icon {icon_cls}">{icon_svg}</div>
                <div class="file-info">
                    <div class="file-name" title="{name_html}">{name_html}</div>
                    <div class="file-meta">
                        <span>ขนาด: {size_mb:.4f} MB</span>
                        <span>แก้ไขล่าสุด: {mod_display}</span>
                    </div>
                </div>
            </div>
"""
                if is_img:
                    html_template += f"""
            <div class="thumbnail-container" onclick="openPreview('{download_path_html}', {name_js})">
                <img class="thumbnail-img" src="{download_path_html}" alt="{name_html}">
            </div>
"""
                html_template += f"""
            <div class="file-actions">
"""
                if is_img:
                    html_template += f"""<button class="action-btn btn-view" onclick="openPreview('{download_path_html}', {name_js})">👁️ ดูภาพ</button>"""
                else:
                    html_template += f"""<a class="action-btn btn-view" href="{download_path_html}" target="_blank" rel="noopener noreferrer">👁️ ดูข้อมูล</a>"""
                html_template += f"""
                <a class="action-btn btn-download" href="{download_path_html}" download>📥 ดาวน์โหลด</a>
                <button class="action-btn btn-delete" onclick="deleteFile({name_js})">🗑️ ลบ</button>
            </div>
        </div>
"""

        html_template += """
    </div>
</div>
<div id="previewModal" class="modal" onclick="closePreview()">
    <div class="modal-content" onclick="event.stopPropagation()">
        <div class="modal-header">
            <span class="modal-title" id="modalTitle">ดูตัวอย่างภาพ</span>
            <span class="close-btn" onclick="closePreview()">&times;</span>
        </div>
        <div class="modal-body"><img id="modalImg" src="" alt="Preview"></div>
    </div>
</div>
<script>
    function filterFiles() {
        const q = document.getElementById('searchInput').value.toLowerCase();
        for (const card of document.getElementsByClassName('file-card')) {
            card.style.display = card.getAttribute('data-name').includes(q) ? 'flex' : 'none';
        }
    }
    function openPreview(src, title) {
        document.getElementById('modalImg').src = src;
        document.getElementById('modalTitle').textContent = title;
        document.getElementById('previewModal').style.display = 'flex';
    }
    function closePreview() {
        document.getElementById('previewModal').style.display = 'none';
    }
    async function deleteFile(fileName) {
        if (!confirm(`คุณแน่ใจหรือไม่ว่าต้องการลบไฟล์ "${fileName}"?`)) return;
        try {
            const res = await fetch(`/files/${encodeURIComponent(fileName)}`, {
                method: 'DELETE', credentials: 'include'
            });
            const data = await res.json();
            if (res.ok && data.success) { alert('ลบไฟล์เรียบร้อยแล้ว!'); window.location.reload(); }
            else if (res.status === 401) { alert('Session หมดอายุ'); window.location.href = '/files'; }
            else { alert('เกิดข้อผิดพลาด: ' + (data.detail || data.message)); }
        } catch (err) { alert('เกิดข้อผิดพลาด: ' + err.message); }
    }
    async function triggerCleanup() {
        if (!confirm('คุณแน่ใจหรือไม่ว่าต้องการสั่งล้างไฟล์ขยะทั้งหมด?')) return;
        try {
            const res = await fetch('/internal/cleanup', { method: 'POST', credentials: 'include' });
            const data = await res.json();
            if (data.success) { alert('ล้างไฟล์ขยะเรียบร้อยแล้ว!'); window.location.reload(); }
            else if (res.status === 401) { alert('Session หมดอายุ'); window.location.href = '/files'; }
            else { alert('เกิดข้อผิดพลาด: ' + data.message); }
        } catch (err) { alert('เกิดข้อผิดพลาด: ' + err.message); }
    }
</script>
</body>
</html>
"""
        return HTMLResponse(content=html_template)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error listing files: {e}", exc_info=True)
        return {"success": False, "message": f"Failed to list files: {str(e)}"}


@app.delete("/files/{file_name}")
async def delete_file(
    file_name: str,
    x_cleanup_token: Optional[str] = Header(None),
    token: Optional[str] = None,
    sc_report_session: Optional[str] = Cookie(None),
):
    if not _is_authenticated_for_files(x_cleanup_token, token, sc_report_session):
        raise HTTPException(status_code=401, detail="Unauthorized")

    safe_name = _safe_report_filename(file_name)
    delete_gcs_file(safe_name)

    return {
        "success": True,
        "message": f"ลบไฟล์ {safe_name} เรียบร้อยแล้ว",
    }


@app.post("/internal/cleanup")
async def trigger_cleanup(
    x_cleanup_token: Optional[str] = Header(None),
    sc_report_session: Optional[str] = Cookie(None),
):
    if CLEANUP_SECRET:
        if not _is_authenticated_for_files(x_cleanup_token, None, sc_report_session):
            raise HTTPException(status_code=401, detail="Unauthorized")
    logger.info("Triggering manual report cleanup via API...")
    deleted = await asyncio.to_thread(cleanup_gcs_files)
    return {"success": True, "message": f"Cleanup completed — {deleted} file(s) removed."}


# ============================================================
# SECTION 23 — MOUNT MCP & ENTRYPOINT
# ============================================================
app.mount("", mcp_asgi_app)

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8080)