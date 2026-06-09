import os
import re
from contextlib import asynccontextmanager
from contextvars import ContextVar
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from fastapi import FastAPI, Request
from fastapi.responses import FileResponse, HTMLResponse
from pydantic import BaseModel
from google.cloud import bigquery
from mcp.server.fastmcp import FastMCP
from fastapi.middleware.cors import CORSMiddleware

# --------------------------------------------------------
# Config
# --------------------------------------------------------
PROJECT_ID = "sc-ai-uat"
DATASET_ID = "SCReport"
AUTH_TABLE = "AuthenByMenu"
DEFAULT_USER_EMAIL = ""
PUBLIC_BASE_URL = os.environ.get(
    "PUBLIC_BASE_URL",
    "https://sc-report-866803019306.asia-southeast3.run.app",
)

REPORT_DIR = Path("/tmp/reports").resolve()
REPORT_DIR.mkdir(parents=True, exist_ok=True)

bq_client = bigquery.Client(project=PROJECT_ID)

CURRENT_REQUEST_USER_EMAIL: ContextVar[Optional[str]] = ContextVar(
    "current_request_user_email",
    default=None,
)

CURRENT_REQUEST_BASE_URL: ContextVar[Optional[str]] = ContextVar(
    "current_request_base_url",
    default=None,
)

# --------------------------------------------------------
# Models
# --------------------------------------------------------
class GenerateReportRequest(BaseModel):
    table_id: str
    user_email: Optional[str] = None
    limit: int = 2000
    condition: Optional[str] = None

# --------------------------------------------------------
# Mapping
# --------------------------------------------------------
TABLE_TO_REPORT_NAMES: dict[str, list[str]] = {
    "vrptexpension": [
        "รายงานตรวจสอบการจ่ายเงิน (ต้นทุน)",
        "รายงานตรวจสอบการจ่ายเงิน(ต้นทุน)",
        "ทะเบียนคุมการเบิกจ่าย (ต้นทุน)",
        "ทะเบียนคุมการเบิกจ่าย(ต้นทุน)",
    ],
    "vrptexpensionexpmodule": [
        "รายงานตรวจสอบการจ่ายเงิน (ค่าใช้จ่าย)",
        "รายงานตรวจสอบการจ่ายเงิน(ค่าใช้จ่าย)",
        "ทะเบียนคุมการเบิกจ่าย (ค่าใช้จ่าย)",
        "ทะเบียนคุมการเบิกจ่าย(ค่าใช้จ่าย)",
    ],
}

TABLE_TO_BQ_TABLE_NAME: dict[str, str] = {
    "vrptexpension": "VRptExpension",
    "vrptexpensionexpmodule": "VRptExpensionExpModule",
}

# --------------------------------------------------------
# MCP Server — ต้อง init ก่อน FastAPI เพื่อใช้ใน lifespan
# --------------------------------------------------------
mcp = FastMCP(
    "sc-report-server",
    stateless_http=True,
    json_response=True,
    host="0.0.0.0",
)

# --------------------------------------------------------
# ASGI sub-app
# --------------------------------------------------------
mcp_asgi_app = mcp.streamable_http_app()

# --------------------------------------------------------
# FastAPI App — lifespan เรียก session_manager.run()
# --------------------------------------------------------
@asynccontextmanager
async def lifespan(app: FastAPI):
    async with mcp.session_manager.run():
        yield

app = FastAPI(
    title="SC Report MCP API & Server (Email Only Version)",
    description="API & MCP Server สำหรับตรวจสอบสิทธิ์ (ผ่าน Email) และสร้าง Excel Report จาก BigQuery",
    version="1.0.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --------------------------------------------------------
# Basic Routes
# --------------------------------------------------------
@app.get("/")
def read_root():
    return {"message": "Welcome to SC Report API. MCP server is active at /mcp"}

@app.get("/health")
def health():
    return {"status": "ok", "service": "sc-report-api-mcp"}

# --------------------------------------------------------
# Middleware — จับ current-user หรือ x-user-email header
# --------------------------------------------------------
@app.middleware("http")
async def capture_user_email_header(request: Request, call_next):
    email_val = None
    header_checked = []
    
    for header_name in ["x-user-email", "current-user", "current_user"]:
        val = request.headers.get(header_name)
        header_checked.append(f"{header_name}={repr(val)}")
        if val:
            email_val = val
            break
    
    print(f"📨 [Middleware] {request.method} {request.url.path}")
    print(f"   Headers checked: {', '.join(header_checked)}")
    print(f"   Resolved email: {repr(email_val)}")
    
    token_email = CURRENT_REQUEST_USER_EMAIL.set(email_val)
    base_url = str(request.base_url).rstrip("/")
    token_base = CURRENT_REQUEST_BASE_URL.set(base_url)

    try:
        response = await call_next(request)
    finally:
        CURRENT_REQUEST_USER_EMAIL.reset(token_email)
        CURRENT_REQUEST_BASE_URL.reset(token_base)
    return response

# --------------------------------------------------------
# Helpers
# --------------------------------------------------------
def clean_table_id(raw: str) -> str:
    return raw.strip().strip("`").split(".")[-1].strip()

def validate_table_id(table_id: str) -> bool:
    return bool(re.match(r"^[A-Za-z0-9_]+$", table_id))

def is_valid_user(user: Optional[str]) -> bool:
    if not user:
        return False
    user_str = user.strip()
    if not user_str:
        return False
    if "${" in user_str or "}" in user_str or "{{" in user_str:
        return False
    return True

def resolve_user_email(provided_email: Optional[str]) -> str:
    print(f"\n🔧 [resolve_user_email] START")
    print(f"   📥 provided_email: {repr(provided_email)}")
    
    context_email = CURRENT_REQUEST_USER_EMAIL.get()
    print(f"   📨 CURRENT_REQUEST_USER_EMAIL (from middleware): {repr(context_email)}")
    
    if is_valid_user(provided_email):
        print(f"   ✅ Step 1: provided_email ถูกต้อง")
        return provided_email
    
    if provided_email is not None and provided_email != "":
        if is_valid_user(context_email):
            return context_email
    
    if provided_email is None:
        if is_valid_user(context_email):
            return context_email
    
    if DEFAULT_USER_EMAIL:
        return DEFAULT_USER_EMAIL
    
    return ""

def check_user_permission(email: str, table_id: str) -> Optional[str]:
    clean = clean_table_id(table_id).lower()
    report_names = TABLE_TO_REPORT_NAMES.get(clean)

    if not report_names:
        return None

    try:
        placeholders = ", ".join([f"@name{i}" for i in range(len(report_names))])
        query = f"""
            SELECT ReportName
            FROM `{PROJECT_ID}.{DATASET_ID}.{AUTH_TABLE}`
            WHERE LOWER(TRIM(Email)) = LOWER(TRIM(@email))
            AND TRIM(ReportName) IN ({placeholders})
            LIMIT 1
        """
        params = [bigquery.ScalarQueryParameter("email", "STRING", email.strip())]
        for i, name in enumerate(report_names):
            params.append(bigquery.ScalarQueryParameter(f"name{i}", "STRING", name))

        job_config = bigquery.QueryJobConfig(query_parameters=params)
        job = bq_client.query(query, job_config=job_config)
        results = list(job)

        if results:
            return results[0]["ReportName"]
        return None
    except Exception as e:
        print(f"🚨 [Auth Error] {str(e)}")
        raise e

def fetch_and_generate_excel(table_id: str, limit: int = 2000, condition: Optional[str] = None) -> Optional[int]:
    clean = clean_table_id(table_id).lower()
    bq_table_name = TABLE_TO_BQ_TABLE_NAME.get(clean, table_id)

    where_clause = ""
    if condition:
        clean_condition = condition.replace(";", "")
        where_clause = f" WHERE {clean_condition}"

    query = f"SELECT * FROM `{PROJECT_ID}.{DATASET_ID}.{bq_table_name}`{where_clause} LIMIT {limit}"

    job = bq_client.query(query)
    result = job.result()
    schema = result.schema
    results = list(result)
    if not results:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    headers = [field.name for field in schema]
    ws.append(headers)

    for row in results:
        row_values = []
        for field in schema:
            val = row[field.name]
            if hasattr(val, "tzinfo") and val.tzinfo is not None:
                val = val.replace(tzinfo=None)
            row_values.append(val)
        ws.append(row_values)

    file_path = REPORT_DIR / f"Report_{table_id}.xlsx"
    wb.save(file_path)
    return len(results)

# --------------------------------------------------------
# MCP Tools
# --------------------------------------------------------
@mcp.tool(
    name="generate_excel_report",
    description=(
        "สร้างไฟล์ Excel จากรายงานใน BigQuery และส่งลิงก์ดาวน์โหลด\n"
        "table_id ที่รองรับ: vrptexpension, vrptexpensionexpmodule"
    ),
)
def mcp_generate_excel_report(
    table_id: str,
    user_email: Optional[str] = None,
    limit: int = 2000,
    condition: Optional[str] = None,
) -> str:
    email_to_use = resolve_user_email(user_email)

    if not validate_table_id(table_id):
        return "❌ table_id ไม่ถูกต้อง"

    report_name = check_user_permission(email_to_use, table_id)
    if not report_name:
        return f"🙏 ขออภัยในความไม่สะดวกครับคุณ {email_to_use} คุณยังไม่มีสิทธิ์เข้าถึงรายงานนี้"

    try:
        row_count = fetch_and_generate_excel(table_id, limit=limit, condition=condition)
        if row_count is None:
            return f"❌ ไม่พบข้อมูลในรายงาน {report_name}"

        file_name = f"Report_{table_id}.xlsx"
        base_url = CURRENT_REQUEST_BASE_URL.get() or PUBLIC_BASE_URL
        download_url = f"{base_url.rstrip('/')}/download/{file_name}"

        return (
            f"📊 ระบบได้จัดเตรียม **รายงาน{report_name}** ({row_count} แถว) เรียบร้อยแล้วครับ\n\n"
            "📥 **คลิกลิงก์ด้านล่างเพื่อดาวน์โหลด (แท็บจะปิดตัวลงเมื่อดาวน์โหลดเสร็จสิ้น):**\n"
            f"🔗 {download_url}"
        )
    except Exception as e:
        return f"🚨 เกิดข้อผิดพลาด: {str(e)}"

# --------------------------------------------------------
# REST API
# --------------------------------------------------------
@app.post("/generate_excel_report")
def generate_excel_report(request: GenerateReportRequest, http_request: Request):
    tid = clean_table_id(request.table_id)
    raw_email = (
        http_request.headers.get("x-user-email")
        or http_request.headers.get("current-user")
        or http_request.headers.get("current_user")
        or request.user_email
        or ""
    )
    email_val = raw_email if is_valid_user(raw_email) else DEFAULT_USER_EMAIL

    if not validate_table_id(tid):
        return {"success": False, "message": f"❌ table_id '{tid}' ไม่ถูกต้อง"}

    report_name = check_user_permission(email_val, tid)
    if not report_name:
        return {"success": False, "message": "🙏 ขออภัย คุณยังไม่มีสิทธิ์เข้าถึงรายงานนี้"}

    try:
        row_count = fetch_and_generate_excel(tid, limit=request.limit, condition=request.condition)
        if row_count is None:
            return {"success": False, "message": "❌ ไม่พบข้อมูล"}

        file_name = f"Report_{tid}.xlsx"
        base_url = CURRENT_REQUEST_BASE_URL.get() or PUBLIC_BASE_URL
        download_url = f"{base_url.rstrip('/')}/download/{file_name}"

        return {
            "success": True,
            "message": f"✅ จัดเตรียม **รายงาน{report_name}** สำเร็จ\n🔗 {download_url}",
        }
    except Exception as e:
        return {"success": False, "message": f"🚨 เกิดข้อผิดพลาด: {str(e)}"}

# --------------------------------------------------------
# Download Excel
# --------------------------------------------------------
@app.get("/download/{file_name}")
def download_report(file_name: str, direct: bool = False):
    if not file_name.endswith(".xlsx"):
        return {"success": False, "message": "❌ อนุญาตเฉพาะไฟล์ .xlsx เท่านั้น"}

    file_path = (REPORT_DIR / file_name).resolve()

    if not file_path.is_relative_to(REPORT_DIR):
        return {"success": False, "message": "❌ ไม่มีสิทธิ์เข้าถึงไฟล์นอกไดเรกทอรีที่กำหนด"}

    if not file_path.exists():
        return {"success": False, "message": "❌ ไม่พบไฟล์ที่ต้องการดาวน์โหลด"}

    if direct:
        return FileResponse(
            path=file_path,
            filename=file_name,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Downloading Report...</title>
    <style>
        body {{
            font-family: system-ui, -apple-system, BlinkMacSystemFont, sans-serif;
            background-color: #0d0f12;
            color: #e2e8f0;
            display: flex;
            justify-content: center;
            align-items: center;
            height: 100vh;
            margin: 0;
        }}
        .loader-container {{
            text-align: center;
        }}
        .spinner {{
            width: 32px;
            height: 32px;
            border: 3px solid rgba(255, 255, 255, 0.1);
            border-top-color: #cda34f;
            border-radius: 50%;
            animation: spin 0.8s linear infinite;
            margin: 0 auto 16px;
        }}
        @keyframes spin {{
            to {{ transform: rotate(360deg); }}
        }}
    </style>
</head>
<body>
    <div class="loader-container">
        <div class="spinner"></div>
        <p style="font-size: 15px; letter-spacing: 0.5px;">กำลังดาวน์โหลดไฟล์และจะปิดหน้าต่างนี้อัตโนมัติ...</p>
        <p style="font-size: 12px; color: #94a3b8; margin-top: 5px;">หากแท็บไม่ปิดเอง คุณสามารถกดปิดหน้าต่างนี้ได้ทันที</p>
    </div>

    <script>
        const downloadUrl = window.location.pathname + "?direct=true";
        const iframe = document.createElement('iframe');
        iframe.style.display = 'none';
        iframe.src = downloadUrl;
        document.body.appendChild(iframe);

        setTimeout(() => {{
            window.close();
        }}, 1000);
    </script>
</body>
</html>"""
    
    return HTMLResponse(content=html_content)

# --------------------------------------------------------
# DEBUG
# --------------------------------------------------------
@app.post("/debug/check_permission")
def debug_check_permission(request: GenerateReportRequest, http_request: Request):
    raw_email = (
        http_request.headers.get("x-user-email")
        or http_request.headers.get("current-user")
        or http_request.headers.get("current_user")
        or request.user_email
        or ""
    )
    email_val = raw_email if is_valid_user(raw_email) else DEFAULT_USER_EMAIL
    clean = clean_table_id(request.table_id)
    report_names = TABLE_TO_REPORT_NAMES.get(clean.lower())

    error_msg = None
    result = None
    try:
        result = check_user_permission(email_val, request.table_id)
    except Exception as e:
        error_msg = str(e)

    return {
        "received_raw_email": raw_email,
        "resolved_email": email_val,
        "is_valid_email": is_valid_user(raw_email),
        "received_table_id": request.table_id,
        "cleaned_table_id": clean.lower(),
        "expected_report_names": report_names,
        "permission_result": result,
        "has_permission": result is not None,
        "error": error_msg,
    }

app.mount("", mcp_asgi_app)

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8080)
